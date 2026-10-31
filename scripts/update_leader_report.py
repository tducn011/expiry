import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from copy import copy

def update_excel():
    excel_path = './Leader/BAO_CAO_TIEN_DO_EXPIRY_GITHUB.xlsx'
    wb = openpyxl.load_workbook(excel_path)
    
    # Border & Fill styles
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    fill_white = PatternFill(fill_type='solid', start_color='FFFFFF', end_color='FFFFFF')
    fill_total = PatternFill(fill_type='solid', start_color='F2F4F7', end_color='F2F4F7')
    
    # ----------------------------------------------------
    # 1. SHEET 2: 📌 Toàn bộ Issues (80) -> 📌 Toàn bộ Issues (93)
    # ----------------------------------------------------
    ws_issues = wb['📌 Toàn bộ Issues (80)']
    ws_issues.title = '📌 Toàn bộ Issues (93)'
    ws_issues['A1'].value = 'DANH SÁCH CHI TIẾT TOÀN BỘ 93 ISSUES ĐÃ GIẢI QUYẾT TRÊN GITHUB'
    ws_issues['A2'].value = 'Tất cả các issues đều đã được hoàn thành 100% (State: CLOSED) | Đồng bộ toàn diện từ GitHub Issues ngày 31/10/2026'
    
    new_issues_data = [
        (81, '#164', 'CSDL & Backend API', '[DB-03] Execute Migration 005_add_category on MySQL and Add Category Filter to /food-entries API', 'Backend API', 'kiet00394-collab', 'enhancement, backend', 'Thực thi migration 005 bổ sung cột category có ràng buộc CHECK constraint 7 nhóm thực phẩm và cập nhật bộ lọc category cho endpoint GET /food-entries API.', 'Hoàn thành', 1.0),
        (82, '#165', 'Frontend & Điều phối', '[FE-14] Wire Session Coordinator and API Live Mode into Storage & Insight Views', 'Frontend Web', 'kiet00394-collab', 'enhancement, frontend', 'Đấu nối Session Coordinator và cờ useDemo/live mode vào màn hình Nơi lưu trữ (StorageView) và Insight (InsightView) để truy vấn trực tiếp CSDL MySQL.', 'Hoàn thành', 1.0),
        (83, '#166', 'Frontend & Điều hướng', '[FE-15] Synchronize Browser History & Deep Linking for Storage and Insight Views', 'Frontend Web', 'kiet00394-collab', 'enhancement, frontend', 'Đồng bộ URL hash/history (pushState/replaceState) cho các màn hình Storage và Insight, hỗ trợ thao tác back/forward mượt mà trên trình duyệt.', 'Hoàn thành', 1.0),
        (84, '#167', 'AI & Quét Barcode', '[SCAN-02] Align Barcode and OCR Scan Flow with publicScanResult Contract', 'AI OCR / Scanner', 'kiet00394-collab', 'enhancement, backend, frontend', 'Chuẩn hóa luồng quét mã vạch và nhận diện OCR theo đúng đặc tả hợp đồng dữ liệu publicScanResult giữa frontend và backend API.', 'Hoàn thành', 1.0),
        (85, '#168', 'Tối ưu Hiệu năng', '[PERF-08] Push LIMIT and OFFSET Pagination into SQL Query for Stock Movement History', 'Backend API', 'kiet00394-collab', 'enhancement, backend', 'Đẩy tham số phân trang LIMIT và OFFSET trực tiếp xuống câu lệnh SQL truy vấn MySQL thay vì cắt mảng trên bộ nhớ ứng dụng.', 'Hoàn thành', 1.0),
        (86, '#169', 'Hạ tầng Build & Tooling', '[BUILD-01] Resolve Vite 8 Configuration Deprecation Warnings', 'DevOps / Tooling', 'kiet00394-collab', 'enhancement, frontend', 'Cập nhật cấu hình vite.config.ts loại bỏ các cảnh báo không tương thích (deprecation warnings) trong phiên bản Vite 8.', 'Hoàn thành', 1.0),
        (87, '#170', 'Thị giác AI & Bảo mật', '[BUG-01] Python Scan OCR Service Health Route Accessible Without Credentials', 'Backend API / AI', 'UncleYelloow', 'bug, backend', 'Loại bỏ route scan healthcheck trùng lặp và mở public endpoint kiểm tra sức khỏe dịch vụ OCR không đòi hỏi token xác thực.', 'Hoàn thành', 1.0),
        (88, '#171', 'CSDL & Tính Toàn Vẹn', '[BUG-02] Resolve Migration Checksum Drift in manage.mjs and Ledger Invariant Mismatch in expiry_dev', 'CSDL MySQL', 'UncleYelloow', 'bug, backend', 'Khắc phục sai lệch mã băm checksum trong manage.mjs và sửa lỗi bất biến số dư sổ cái kho ledger trong database expiry_dev.', 'Hoàn thành', 1.0),
        (89, '#181', 'Bảo mật & Xác thực', '[BUG-03] Frontend AuthForm does not call backend auth API endpoints', 'Frontend / Auth', 'UncleYelloow', 'bug, frontend, security', 'Đấu nối form đăng nhập/đăng ký AuthForm trực tiếp vào các endpoint API backend (/api/auth/login, /api/auth/register, /api/auth/me).', 'Hoàn thành', 1.0),
        (90, '#182', 'Frontend & Kho hàng', '[BUG-04] App.tsx draft submission does not persist food entries to MySQL API', 'Frontend Web', 'UncleYelloow', 'bug, backend, frontend', 'Kết nối hàm lưu nháp handleSaveDraft trong App.tsx vào API backend để lưu thực phẩm vào CSDL MySQL khi chạy chế độ Live.', 'Hoàn thành', 1.0),
        (91, '#183', 'Quản lý Kho & Sổ cái', '[BUG-05] Detail view stock movements bypass database movement ledger', 'Frontend Web', 'UncleYelloow', 'bug, backend, frontend', 'Đấu nối các thao tác dùng, bỏ và kiểm kê số lượng trong Detail.tsx vào backend API thay vì chỉ lưu tạm trên frontend state.', 'Hoàn thành', 1.0),
        (92, '#184', 'Thị giác AI & Camera', '[BUG-06] Camera Scanner uses hardcoded timeout mock instead of calling Python Scan OCR Service', 'Frontend / AI', 'UncleYelloow', 'bug, backend, frontend', 'Thay thế mock timer bằng cuộc gọi API multipart upload thực tế tới dịch vụ Python OCR để nhận diện hạn sử dụng từ ảnh camera.', 'Hoàn thành', 1.0),
        (93, '#186', 'Vòng đời Thùng rác', '[BUG-07] Trash lifecycle soft-delete and restore actions bypass backend API', 'Frontend Web', 'UncleYelloow', 'bug, backend, frontend', 'Kết nối thao tác xóa mềm vào thùng rác và khôi phục món ăn trong Trash.tsx vào backend API MySQL (/api/food-entries/:id/trash, /restore).', 'Hoàn thành', 1.0)
    ]
    
    # Overwrite old total row (row 85) with first new item, then append the rest
    for i, item in enumerate(new_issues_data):
        row_idx = 85 + i
        for col_idx, val in enumerate(item, start=1):
            cell = ws_issues.cell(row_idx, col_idx, val)
            cell.font = Font(name='Calibri', size=10, bold=(col_idx in (1, 2, 9)))
            cell.border = thin_border
            cell.fill = fill_white
            if col_idx in (1, 2, 6, 9):
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx == 10:
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = '0.0%'
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=(col_idx in (4, 8)))
                
    # Add new Total row at row 98
    total_row = 85 + len(new_issues_data) # 98
    ws_issues.cell(total_row, 1, 'TỔNG').font = Font(name='Calibri', size=10, bold=True)
    ws_issues.cell(total_row, 2, f'=COUNTA(B5:B{total_row-1})').font = Font(name='Calibri', size=10, bold=True)
    ws_issues.cell(total_row, 3, '15 EPICs + 78 Tasks & Bugs').font = Font(name='Calibri', size=10, bold=True)
    ws_issues.cell(total_row, 4, f'Toàn bộ {total_row-5} Issues đã đóng').font = Font(name='Calibri', size=10, bold=True)
    ws_issues.cell(total_row, 8, '100% mục tiêu kỹ thuật đã hoàn thành và kiểm chứng').font = Font(name='Calibri', size=10, bold=True)
    ws_issues.cell(total_row, 9, f'=COUNTIF(I5:I{total_row-1}, "Hoàn thành")').font = Font(name='Calibri', size=10, bold=True)
    c10 = ws_issues.cell(total_row, 10, f'=COUNTIF(I5:I{total_row-1}, "Hoàn thành")/COUNTA(B5:B{total_row-1})')
    c10.font = Font(name='Calibri', size=10, bold=True)
    c10.number_format = '0.0%'
    for c in range(1, 11):
        cell = ws_issues.cell(total_row, c)
        cell.border = thin_border
        cell.fill = fill_total

    # ----------------------------------------------------
    # 2. SHEET 3: 🔀 Merged PRs develop (67) -> 🔀 Merged PRs develop (80)
    # ----------------------------------------------------
    ws_prs = wb['🔀 Merged PRs develop (67)']
    ws_prs.title = '🔀 Merged PRs develop (80)'
    ws_prs['A1'].value = 'DANH SÁCH 80 PULL REQUESTS ĐÃ MERGE VÀO NHÁNH DEVELOP'
    ws_prs['A2'].value = 'Tất cả 80 PRs đều vượt qua kiểm tra chất lượng CI/CD tự động trước khi merge vào develop | Cập nhật: 31/10/2026'

    new_prs_data = [
        (68, '#172', 'feat(core): implement category taxonomy, storage hub view, insight analytics, and active tunnel routing', 'tducn110 (Leader)', 'feat/category-storage-insight-tunnel', '2026-10-28', '-', 'Triển khai phân loại 7 nhóm thực phẩm, màn hình Storage Hub trực quan hóa 4 vị trí lưu trữ, màn hình Insight phân tích thống kê và định tuyến tunnel URL.', 'Hoàn thành', 1.0),
        (69, '#174', 'style(web): format code with oxfmt and fix type definitions', 'tducn110 (Leader)', 'fix/oxfmt-formatting-gates', '2026-10-28', '-', 'Định dạng lại toàn bộ mã nguồn frontend web bằng oxfmt, sửa các lỗi cú pháp TypeScript interface và đảm bảo CI quality gate đạt 100%.', 'Hoàn thành', 1.0),
        (70, '#177', 'feat(web): add collapsible dashboard navigation and clean oxfmt type interfaces', 'tducn110 (Leader)', 'fix/collapsible-dashboard-and-oxfmt-cleanup', '2026-10-28', '-', 'Thêm thanh điều hướng Sidebar có thể thu gọn (collapsible), dọn dẹp các khai báo thừa và chuẩn hóa tương thích oxfmt.', 'Hoàn thành', 1.0),
        (71, '#178', 'feat(web): integrate @vercel/analytics and @vercel/speed-insights', 'tducn110 (Leader)', 'feat/vercel-analytics-speed-insights', '2026-10-28', '-', 'Tích hợp thư viện @vercel/analytics và @vercel/speed-insights cho ứng dụng React SPA để đo lường hiệu năng Core Web Vitals và lượt truy cập thực tế.', 'Hoàn thành', 1.0),
        (72, '#180', 'fix(web): resolve analytics empty states, smart advice icons, and storage view icon consistency', 'tducn110 (Leader)', 'fix/ui-ux-analytics-and-icon-consistency', '2026-10-28', '-', 'Sửa trạng thái trống (empty state) cho màn hình analytics, chuẩn hóa icon lời khuyên thông minh và đồng bộ icon khu vực lưu trữ trên StorageView.', 'Hoàn thành', 1.0),
        (73, '#187', 'fix: resolve bugs BUG-01 to BUG-07', 'UncleYelloow', 'fix/resolve-bugs-01-07', '2026-10-29', '#170, #171, #181, #182, #183, #184, #186', 'Khắc phục triệt để toàn bộ 7 lỗi backend integration phát hiện trên Live server: kết nối auth API, lưu thực phẩm, ghi nhận sổ cái biến động kho, gọi dịch vụ OCR camera và khôi phục thùng rác.', 'Hoàn thành', 1.0),
        (74, '#188', 'merge: reconcile main ancestry into develop', 'UncleYelloow', 'reconcile/sync-main-into-develop', '2026-10-29', '-', 'Đồng bộ nhánh main vào nhánh develop, hợp nhất lịch sử commit và bảo đảm develop kế thừa đầy đủ mọi commit của main.', 'Hoàn thành', 1.0),
        (75, '#189', 'fix(web): resolve Vite 8 configuration deprecation warnings (#169)', 'kiet00394-collab', 'fix/issue-169-vite8-deprecation-warnings', '2026-10-29', '#169', 'Khắc phục các cảnh báo deprecation trong file cấu hình vite.config.ts khi chạy với Vite 8.', 'Hoàn thành', 1.0),
        (76, '#190', 'perf(api): push pagination parameters into SQL query for stock movements (#168)', 'kiet00394-collab', 'fix/issue-168-push-pagination-sql-stock-movements', '2026-10-29', '#168', 'Tối ưu hóa hiệu năng API lịch sử kho bằng cách đẩy tham số LIMIT và OFFSET vào câu lệnh SQL truy vấn MySQL.', 'Hoàn thành', 1.0),
        (77, '#191', 'feat(db,api): execute migration 005 and add category filter to food-entries (#164)', 'kiet00394-collab', 'fix/issue-164-migration-005-category-filter', '2026-10-29', '#164', 'Chạy migration 005_add_category trên CSDL MySQL và cập nhật bộ lọc danh mục category cho endpoint GET /food-entries.', 'Hoàn thành', 1.0),
        (78, '#192', 'feat(scan,ui): align barcode and OCR scan flow with publicScanResult (#167)', 'kiet00394-collab', 'fix/issue-167-scan-flow-public-contract', '2026-10-29', '#167', 'Đồng bộ luồng quét mã vạch và nhận diện OCR theo đúng hợp đồng dữ liệu publicScanResult giữa frontend và backend.', 'Hoàn thành', 1.0),
        (79, '#193', 'feat(web): synchronize browser history and deep linking for storage and insight views (#166)', 'kiet00394-collab', 'fix/issue-166-browser-history-deep-linking', '2026-10-29', '#166', 'Đồng bộ lịch sử trình duyệt và hỗ trợ liên kết sâu (deep linking) cho 2 màn hình mới StorageView và InsightView.', 'Hoàn thành', 1.0),
        (80, '#194', 'feat(web): wire session coordinator and api live mode into storage and insight (#165)', 'kiet00394-collab', 'fix/issue-165-wire-session-coordinator-storage-insight', '2026-10-29', '#165', 'Đấu nối Session Coordinator và chế độ gọi API thực (Live Mode) cho màn hình Nơi lưu trữ và Insight.', 'Hoàn thành', 1.0)
    ]

    for i, item in enumerate(new_prs_data):
        row_idx = 72 + i
        for col_idx, val in enumerate(item, start=1):
            cell = ws_prs.cell(row_idx, col_idx, val)
            cell.font = Font(name='Calibri', size=10, bold=(col_idx in (1, 2, 9)))
            cell.border = thin_border
            cell.fill = fill_white
            if col_idx in (1, 2, 4, 6, 9):
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx == 10:
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = '0.0%'
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=(col_idx in (3, 8)))

    # Total row for PRs at row 85
    pr_total_row = 72 + len(new_prs_data) # 85
    ws_prs.cell(pr_total_row, 1, 'TỔNG').font = Font(name='Calibri', size=10, bold=True)
    ws_prs.cell(pr_total_row, 2, f'=COUNTA(B5:B{pr_total_row-1})').font = Font(name='Calibri', size=10, bold=True)
    ws_prs.cell(pr_total_row, 3, f'{pr_total_row-5} PRs merged vào nhánh develop').font = Font(name='Calibri', size=10, bold=True)
    ws_prs.cell(pr_total_row, 8, f'Toàn bộ {pr_total_row-5} PRs đều vượt qua CI/CD và merge thành công').font = Font(name='Calibri', size=10, bold=True)
    ws_prs.cell(pr_total_row, 9, f'=COUNTIF(I5:I{pr_total_row-1}, "Hoàn thành")').font = Font(name='Calibri', size=10, bold=True)
    c10_pr = ws_prs.cell(pr_total_row, 10, f'=COUNTIF(I5:I{pr_total_row-1}, "Hoàn thành")/COUNTA(B5:B{pr_total_row-1})')
    c10_pr.font = Font(name='Calibri', size=10, bold=True)
    c10_pr.number_format = '0.0%'
    for c in range(1, 11):
        cell = ws_prs.cell(pr_total_row, c)
        cell.border = thin_border
        cell.fill = fill_total

    # ----------------------------------------------------
    # 3. SHEET 4: 🚀 Tính năng & Deliverables
    # ----------------------------------------------------
    ws_deliv = wb['🚀 Tính năng & Deliverables']
    new_deliv_data = [
        (27, 'DB-03', 'Cơ Sở Dữ Liệu', 'Lược Đồ Phân Loại Danh Mục (Category Taxonomy)', 'Migration 005 thêm cột category với CHECK constraint 7 nhóm thực phẩm (dairy, meat, seafood, produce, pantry, beverage, other), hỗ trợ bộ lọc API theo category.', 'database/mysql/migrations/005_add_category.sql, apps/api/src/modules/inventory/', 'PR #191, Issue #164', 'Hoàn thành', 1.0),
        (28, 'FE-07', 'Frontend Web', 'Kho Lưu Trữ Storage Hub & Phân Tích Thông Minh Insight', 'Giao diện StorageView hiển thị trực quan thực phẩm theo 4 khu vực lưu trữ; InsightView phân tích thống kê lãng phí, tỉ lệ tiêu thụ và cảnh báo trực quan bằng biểu đồ.', 'apps/web/src/features/StorageView.tsx, InsightView.tsx', 'PR #172, PR #180, PR #194', 'Hoàn thành', 1.0),
        (29, 'FE-08', 'Frontend Web', 'Lịch Sử Trình Duyệt Deep Linking & Session Coordinator', 'Đồng bộ URL hash/history (pushState/replaceState), hỗ trợ nút Back/Forward trình duyệt và điều phối phiên làm việc giữa chế độ Live Mode và Demo Mode.', 'apps/web/src/App.tsx, store/appStore.ts, hooks/useExpiry.ts', 'PR #193, PR #194, Issue #165, #166', 'Hoàn thành', 1.0),
        (30, 'OPS-02', 'DevOps & Đo Lường', 'Tích Hợp Vercel Analytics & Speed Insights', 'Tích hợp bộ đo lường trải nghiệm người dùng thực (Real User Monitoring), theo dõi các chỉ số Core Web Vitals (LCP, FID, CLS) và phân tích lưu lượng người dùng.', 'apps/web/src/main.tsx, package.json', 'PR #178', 'Hoàn thành', 1.0)
    ]
    
    for i, item in enumerate(new_deliv_data):
        row_idx = 31 + i
        for col_idx, val in enumerate(item, start=1):
            cell = ws_deliv.cell(row_idx, col_idx, val)
            cell.font = Font(name='Calibri', size=10, bold=(col_idx in (1, 2, 8)))
            cell.border = thin_border
            cell.fill = fill_white
            if col_idx in (1, 2, 8):
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx == 9:
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = '0.0%'
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=(col_idx in (4, 5)))
                
    deliv_total_row = 31 + len(new_deliv_data) # 35
    ws_deliv.cell(deliv_total_row, 1, 'TỔNG').font = Font(name='Calibri', size=10, bold=True)
    ws_deliv.cell(deliv_total_row, 2, f'=COUNTA(B5:B{deliv_total_row-1})').font = Font(name='Calibri', size=10, bold=True)
    ws_deliv.cell(deliv_total_row, 3, '7 Phân hệ chính').font = Font(name='Calibri', size=10, bold=True)
    ws_deliv.cell(deliv_total_row, 4, f'Toàn bộ {deliv_total_row-5} Tính năng').font = Font(name='Calibri', size=10, bold=True)
    ws_deliv.cell(deliv_total_row, 5, 'Toàn bộ deliverables đã bàn giao và hoạt động ổn định').font = Font(name='Calibri', size=10, bold=True)
    ws_deliv.cell(deliv_total_row, 8, f'=COUNTIF(H5:H{deliv_total_row-1}, "Hoàn thành")').font = Font(name='Calibri', size=10, bold=True)
    c9_del = ws_deliv.cell(deliv_total_row, 9, f'=COUNTIF(H5:H{deliv_total_row-1}, "Hoàn thành")/COUNTA(B5:B{deliv_total_row-1})')
    c9_del.font = Font(name='Calibri', size=10, bold=True)
    c9_del.number_format = '0.0%'
    for c in range(1, 10):
        cell = ws_deliv.cell(deliv_total_row, c)
        cell.border = thin_border
        cell.fill = fill_total

    # ----------------------------------------------------
    # 4. SHEET 5: ⚠️ Tồn đọng & Đề xuất
    # ----------------------------------------------------
    ws_back = wb['⚠️ Tồn đọng & Đề xuất']
    # BACKLOG-06 (row 10) was Offline-First / Live Mode Sync -> Update as Completed in Phase 10 via Session Coordinator
    ws_back.cell(10, 8, 'Hoàn thành (PR #194)')
    ws_back.cell(10, 9, 1.0)

    # ----------------------------------------------------
    # 5. SHEET 1: 📊 Dashboard & Tổng quan
    # ----------------------------------------------------
    ws_dash = wb['📊 Dashboard & Tổng quan']
    ws_dash['A2'].value = 'Smart Expiry Date & Food Inventory Management System | Nguồn dữ liệu: GitHub Repo tducn011/expiry | Cập nhật: 31/10/2026 (sau PR #194 & đồng bộ toàn diện GitHub)'
    ws_dash['B4'].value = '93 / 93'
    ws_dash['D4'].value = '80 PRs'
    ws_dash['F4'].value = '142'
    ws_dash['H4'].value = '122 / 122'
    ws_dash['J4'].value = '100.0%'
    
    # Table 1: Bảng kiểm tra tiến độ theo từng tab
    ws_dash['B9'].value = '📌 Toàn bộ Issues (93)'
    ws_dash['C9'].value = 'Toàn bộ 93 issues từ GitHub (15 EPICs + 78 Tasks chi tiết & Bugs)'
    ws_dash['D9'].value = "='📌 Toàn bộ Issues (93)'!B98"
    ws_dash['E9'].value = "='📌 Toàn bộ Issues (93)'!I98"
    ws_dash['F9'].value = 0
    ws_dash['G9'].value = "='📌 Toàn bộ Issues (93)'!J98"
    ws_dash['H9'].value = 'Hoàn thành 100%'

    ws_dash['B10'].value = '🔀 Merged PRs develop (80)'
    ws_dash['C10'].value = 'Toàn bộ 80 Pull Requests đã merge vào nhánh develop'
    ws_dash['D10'].value = "='🔀 Merged PRs develop (80)'!B85"
    ws_dash['E10'].value = "='🔀 Merged PRs develop (80)'!I85"
    ws_dash['F10'].value = 0
    ws_dash['G10'].value = "='🔀 Merged PRs develop (80)'!J85"
    ws_dash['H10'].value = 'Hoàn thành 100%'

    ws_dash['B11'].value = '🚀 Tính năng & Deliverables'
    ws_dash['C11'].value = 'Bảng kiểm 30 tính năng sản phẩm cụ thể (FE, BE, AI OCR, DB, CI/CD, Design System, Storage & Insight)'
    ws_dash['D11'].value = "='🚀 Tính năng & Deliverables'!B35"
    ws_dash['E11'].value = "='🚀 Tính năng & Deliverables'!H35"
    ws_dash['F11'].value = 0
    ws_dash['G11'].value = "='🚀 Tính năng & Deliverables'!I35"
    ws_dash['H11'].value = 'Đã bàn giao đầy đủ'

    # Table 2: Phân bổ đóng góp thành viên
    ws_dash['D16'].value = 'Kiến trúc Monorepo, Setup CI/CD, MySQL Migrations 001-005, Vercel Reverse Proxy, AI Scan OCR, Polymorphic ApiClient, Signed Release main, Modern Design System, Storage Hub & Insight Analytics, Session Coordinator (PR #161, #172, #174, #177, #178, #180)'
    ws_dash['E16'].value = 25
    ws_dash['F16'].value = 60

    ws_dash['D17'].value = 'Core Domain Kho (Stock Movements, Ledger, Attention Engine, Food Detail, Trash Lifecycle), Manual Capture Form, Quét Barcode GS1, khắc phục triệt để toàn bộ 7 lỗi backend integration (BUG-01 -> BUG-07, PR #187, #188)'
    ws_dash['E17'].value = 17
    ws_dash['F17'].value = 41

    ws_dash['D18'].value = 'Chiến dịch Ponytail (Cắt giảm 1,100+ dòng code), Tối ưu hiệu năng N+1 & Waterfall, Debounce 300ms, Migration 005 CSDL, SQL pagination pushdown, Vite 8 upgrade, Deep Linking browser history (PR #189 -> #194)'
    ws_dash['E18'].value = 38
    ws_dash['F18'].value = 41

    ws_dash['E19'].value = '=SUM(E16:E18)'
    ws_dash['F19'].value = '=SUM(F16:F18)'
    
    # Table 3: Kết quả đột phá
    extra_breakthroughs = [
        ('B27', '6. Tích Hợp AI Scan OCR & Quét Mã Vạch GS1 Đa Nguồn', 'C27', 'Kết nối camera scanner và file upload trực tiếp tới microservice Python YOLO+CTC OCR; tích hợp ZXing quét barcode GS1 và tự động đối soát sản phẩm.'),
        ('B28', '7. Kho Lưu Trữ Storage Hub & Phân Tích Thông Minh Insight', 'C28', 'Bổ sung màn hình Storage Hub phân chia trực quan 4 khu vực lưu trữ; màn hình Insight phân tích thống kê lãng phí, tỉ lệ tiêu thụ và dự báo trực quan bằng biểu đồ.'),
        ('B29', '8. Kết Nối Đồng Bộ Toàn Diện Real API & Session Coordinator', 'C29', 'Xử lý dứt điểm 7 lỗi kết nối live (BUG-01 -> BUG-07), đồng bộ Auth, thêm/sửa món ăn, xuất/nhập kho, phục hồi thùng rác vào CSDL MySQL với Session Coordinator và deep linking URL.')
    ]
    for b_cell, b_val, c_cell, c_val in extra_breakthroughs:
        ws_dash[b_cell].value = b_val
        ws_dash[b_cell].font = Font(name='Calibri', size=11, bold=True)
        ws_dash[b_cell].border = thin_border
        ws_dash[b_cell].fill = fill_white
        ws_dash[b_cell].alignment = Alignment(horizontal='left', vertical='center')
        
        ws_dash[c_cell].value = c_val
        ws_dash[c_cell].font = Font(name='Calibri', size=10, bold=False)
        ws_dash[c_cell].border = thin_border
        ws_dash[c_cell].fill = fill_white
        ws_dash[c_cell].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

    wb.save(excel_path)
    print('Successfully updated', excel_path)

if __name__ == '__main__':
    update_excel()
